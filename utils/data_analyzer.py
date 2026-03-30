"""
Data Analyzer for structured files (Excel / CSV).

When users upload spreadsheets and ask analytical questions
(highest, lowest, average, total, count, etc.), this module
computes the answer directly from the data rather than relying
on text-similarity retrieval.
"""

import os
import re
import csv
import json
from typing import Dict, List, Optional, Any

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from groq import Groq
except ImportError:
    Groq = None


# ── Keyword patterns that signal an analytical question ──────────────────────
AGGREGATE_PATTERNS = [
    (r"\b(highest|maximum|max|most|top|greatest|best)\b", "max"),
    (r"\b(lowest|minimum|min|least|worst|bottom|fewest)\b", "min"),
    (r"\b(average|mean|avg)\b", "avg"),
    (r"\b(total|sum|overall)\b", "sum"),
    (r"\b(count|how many|number of)\b", "count"),
    (r"\b(sort|rank|order|list all)\b", "sort"),
]

# Patterns for filter/conditional queries
FILTER_PATTERNS = [
    # "greater than 80", "above 90", "more than 75", "over 80", "at least 80"
    (r"(greater than|above|more than|over|at least|>=?|exceeds?)\s*(\d+\.?\d*)", "gte"),
    # "less than 80", "below 70", "under 60", "at most 50"
    (r"(less than|below|under|at most|<=?)\s*(\d+\.?\d*)", "lte"),
    # "equal to 80", "exactly 80"
    (r"(equal to|exactly|equals?)\s*(\d+\.?\d*)", "eq"),
    # "between 70 and 90"
    (r"between\s+(\d+\.?\d*)\s*(?:and|to|-)\s*(\d+\.?\d*)", "between"),
]


class StructuredDataStore:
    """Keeps in-memory tables from uploaded Excel / CSV files."""

    def __init__(self):
        # { filename: [ {col: val, …}, … ] }
        self.tables: Dict[str, List[Dict[str, Any]]] = {}
        # { filename: [col_names] }
        self.headers: Dict[str, List[str]] = {}

    # ── Loading ──────────────────────────────────────────────────────────────
    def load_excel(self, file_path: str) -> int:
        """Load all sheets from an Excel file. Returns row count."""
        if openpyxl is None:
            return 0

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        total = 0
        fname = os.path.basename(file_path)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue

            # Auto-detect real header row (skip merged title rows)
            header_idx = self._find_header_row(rows)
            headers = [str(h).strip() if h is not None else f"Col{i}"
                       for i, h in enumerate(rows[header_idx])]

            records = []
            for row in rows[header_idx + 1:]:
                cells = list(row)
                filled = [c for c in cells if c is not None and str(c).strip()]
                if len(filled) < 2:
                    continue
                # Skip rows without a text name (totals / max-marks)
                has_name = any(
                    isinstance(c, str) and len(c.strip()) > 3 and not c.strip().replace('.', '').isdigit()
                    for c in cells
                )
                if not has_name:
                    continue
                record = {}
                for h, cell in zip(headers, cells):
                    record[h] = cell
                records.append(record)

            if records:
                key = f"{fname}::{sheet_name}" if len(wb.sheetnames) > 1 else fname
                self.tables[key] = records
                self.headers[key] = headers
                total += len(records)

        wb.close()
        return total

    @staticmethod
    def _find_header_row(rows) -> int:
        """Find the real header row by looking for keyword matches."""
        kw = {'name', 'no', 'roll', 'sl', 'sno', 'total', 'id',
              'section', 'subject', 'marks', 'grade', 'percentage',
              'attendance', 'date', 'class', 'student'}
        best_idx, best_score = 0, 0
        for i, row in enumerate(rows[:20]):
            cells = [str(c).strip().lower() for c in row if c is not None and str(c).strip()]
            if len(cells) < 3:
                continue
            hits = sum(1 for c in cells if any(k in c for k in kw))
            short = sum(1 for c in cells if len(c) < 30)
            score = hits * 3 + short
            if score > best_score:
                best_score = score
                best_idx = i
        return best_idx

    def load_csv(self, file_path: str) -> int:
        """Load a CSV file. Returns row count."""
        fname = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            records = list(reader)

        if not records:
            return 0

        self.tables[fname] = records
        self.headers[fname] = list(records[0].keys())
        return len(records)

    def clear(self):
        self.tables.clear()
        self.headers.clear()

    @property
    def has_data(self) -> bool:
        return bool(self.tables)

    # ── Analysis ─────────────────────────────────────────────────────────────
    def _query_mentions_specific_entity(self, query: str) -> bool:
        """Check if the query references a specific ID/roll number or known name."""
        # Check for roll number patterns
        if self._ID_PATTERN.search(query) or self._GENERIC_ID.search(query):
            return True
        # Check if any known cell value (name/ID) appears in the query
        q_lower = query.lower()
        for tkey, rows in self.tables.items():
            for row in rows:
                for val in row.values():
                    if val is None:
                        continue
                    val_str = str(val).strip()
                    if len(val_str) >= 3 and val_str.lower() in q_lower:
                        return True
        return False

    def answer_query(self, query: str) -> Optional[str]:
        """
        Try to answer a query by analysing the stored structured data.
        Returns an answer string, or None if the query isn't analytical.
        """
        if not self.has_data:
            return None

        # 0) Try comparison first ("compare X and Y", "who is better X or Y")
        ans = self._try_comparison(query)
        if ans:
            return ans

        # If query mentions a specific student/ID, try row lookup FIRST
        if self._query_mentions_specific_entity(query):
            ans = self._try_row_lookup(query)
            if ans:
                return ans

        # 1) Try filter + count ("how many students have attendance > 80%")
        ans = self._try_filter_query(query)
        if ans:
            return ans

        # 2) Try aggregate (highest, lowest, avg, total, count, rank)
        op = self._detect_operation(query)
        if op is not None:
            table_key, column = self._match_column(query, op)
            if op == "count" and table_key is None:
                table_key = next(iter(self.tables))
                column = None
            if table_key is not None:
                rows = self.tables[table_key]
                result = self._compute(rows, column, op, query)
                if result:
                    return result

        # 3) Try row lookup ONLY if query looks like a person/ID lookup
        #    (not for general knowledge questions about PDF content)
        if self._is_entity_query(query):
            ans = self._try_row_lookup(query)
            if ans:
                return ans

        # 4) Fallback: Use Groq LLM to analyze the data for complex questions
        ans = self._try_llm_analysis(query)
        if ans:
            return ans

        return None

    def _is_entity_query(self, query: str) -> bool:
        """Check if the query is asking about a specific person/ID/record,
        not a general knowledge question."""
        # Has a roll number / ID pattern
        if self._ID_PATTERN.search(query) or self._GENERIC_ID.search(query):
            return True
        # Has a name in ALL CAPS (like student names)
        if re.search(r'\b[A-Z][A-Z ]{4,}\b', query):
            return True
        # Query patterns that suggest a person lookup
        person_patterns = (
            r'\bwho is\b', r'\btell me about\b', r'\bdetails of\b',
            r'\battendance of\b', r'\bmarks of\b', r'\bscore of\b',
        )
        q_lower = query.lower()
        if any(re.search(p, q_lower) for p in person_patterns):
            # But only if the query is short (likely a name lookup, not a concept question)
            # "who is mahesh babu" = name lookup
            # "what is hallucination firewall" = concept question
            words = query.split()
            if len(words) <= 8:
                return True
        return False

    # ── Row Lookup ────────────────────────────────────────────────────────────

    # Patterns that look like IDs / roll numbers (alphanumeric codes)
    _ID_PATTERN = re.compile(r'\b(\d{2}[A-Za-z]{2}\d[A-Za-z]\d{4})\b')  # e.g. 22PA1A0504
    _GENERIC_ID = re.compile(r'\b([A-Z]{2,}\d{3,}[A-Z]*\d*)\b', re.IGNORECASE)  # broader

    def _try_row_lookup(self, query: str) -> Optional[str]:
        """Answer queries like 'How many UHV classes attended by 22PA1A0501?'
        or 'What is the attendance of 22PA1A0504?' or 'Tell me about Alice'.

        If the query mentions a specific column, returns only that value.
        Otherwise returns the full row.
        If the query mentions an ID/roll number that doesn't exist, flags it
        as hallucinated.
        """
        q_lower = query.lower()

        for tkey, rows in self.tables.items():
            headers = self.headers[tkey]
            label_col = self._find_label_column(rows)

            for row in rows:
                # Check every cell value in the row for a match with the query
                matched_id = None
                for col in headers:
                    val = row.get(col)
                    if val is None:
                        continue
                    val_str = str(val).strip()
                    if len(val_str) < 3:
                        continue
                    if val_str.lower() in q_lower:
                        matched_id = val_str
                        break

                if matched_id is None:
                    continue

                # Found the row — now figure out what the user is asking
                name_val = str(row.get(label_col, matched_id)).strip()

                # ── Check if the query is a verification/claim question ────
                # e.g. "is 22PA1A0501 has attendance percentage of 90%"
                claimed_value = self._extract_claimed_value(query)
                asked_cols = self._find_asked_columns(query, headers, tkey)

                if claimed_value is not None and asked_cols:
                    # User is claiming a specific value — verify it
                    for ac in asked_cols:
                        actual = self._to_float(row.get(ac))
                        if actual is not None:
                            if abs(actual - claimed_value) < 0.5:
                                return (
                                    f"Yes, that is correct. The {ac} of {name_val} is {actual}, "
                                    f"which matches the claimed value of {claimed_value}."
                                )
                            else:
                                return (
                                    f"HALLUCINATION DETECTED: No, that is incorrect. "
                                    f"The claimed {ac} of {name_val} is {claimed_value}, "
                                    f"but the actual value is {actual}. "
                                    f"The claim does not match the uploaded data."
                                )
                elif claimed_value is not None:
                    # User claimed a value but no specific column detected — check all numeric columns
                    for h in headers:
                        actual = self._to_float(row.get(h))
                        if actual is not None and abs(actual - claimed_value) < 0.5:
                            return (
                                f"Yes, that is correct. The {h} of {name_val} is {actual}, "
                                f"which matches the claimed value of {claimed_value}."
                            )
                    # No column matched the claimed value
                    # Find the most likely column (e.g. % or total)
                    likely_cols = [h for h in headers if h.strip() in ('%', 'TOTAL', 'Percentage')]
                    if likely_cols:
                        ac = likely_cols[0]
                        actual = self._to_float(row.get(ac))
                        if actual is not None:
                            return (
                                f"HALLUCINATION DETECTED: No, that is incorrect. "
                                f"The claimed value for {name_val} is {claimed_value}, "
                                f"but the actual {ac} is {actual}. "
                                f"The claim does not match the uploaded data."
                            )

                if asked_cols:
                    # Return only the requested fields
                    parts = []
                    for ac in asked_cols:
                        cell = row.get(ac)
                        if cell is not None:
                            parts.append(f"{ac}: {cell}")

                    if len(parts) == 1:
                        col_name, col_val = parts[0].split(": ", 1)
                        return f"The {col_name} of {name_val} is {col_val}."
                    else:
                        return f"For {name_val}:\n" + "\n".join(f"  - {p}" for p in parts)
                else:
                    # No specific column detected — return full row
                    parts = []
                    for h in headers:
                        cell = row.get(h)
                        if cell is not None and str(cell).strip():
                            parts.append(f"{h}: {cell}")
                    return f"Details for {name_val}:\n" + "\n".join(f"  - {p}" for p in parts)

        # ── No row matched — check if the query contains an ID that looks
        #    like it *should* be in the data but isn't (hallucination) ────────
        return self._check_hallucinated_id(query)

    # ── Student Comparison ───────────────────────────────────────────────────
    _COMPARE_PATTERNS = re.compile(
        r'(compare|versus|vs\.?|difference between|who.*(better|higher|more|greater|lower|less|worse))',
        re.IGNORECASE
    )

    def _try_comparison(self, query: str) -> Optional[str]:
        """Handle queries like 'compare 22PA1A0501 and 22PA1A0502' or
        'who has better attendance 22PA1A0501 or 22PA1A0502'."""
        if not self._COMPARE_PATTERNS.search(query):
            return None

        # Find all entity matches (IDs or names) in the query
        matched_rows = []
        q_lower = query.lower()

        for tkey, rows in self.tables.items():
            headers = self.headers[tkey]
            label_col = self._find_label_column(rows)
            for row in rows:
                for col in headers:
                    val = row.get(col)
                    if val is None:
                        continue
                    val_str = str(val).strip()
                    if len(val_str) < 3:
                        continue
                    if val_str.lower() in q_lower:
                        name_val = str(row.get(label_col, val_str)).strip()
                        if not any(r[0] == name_val for r in matched_rows):
                            matched_rows.append((name_val, row, headers, tkey))
                        break

        if len(matched_rows) < 2:
            # Extract all IDs/names mentioned in the query
            requested_ids = self._ID_PATTERN.findall(query)
            requested_ids += self._GENERIC_ID.findall(query)
            # Also check for full names in caps
            requested_ids += re.findall(r'\b([A-Z][A-Z ]{4,})\b', query)

            if len(matched_rows) == 1 and len(requested_ids) >= 2:
                # One student found, one not — partial verification
                found_name = matched_rows[0][0]
                found_row = matched_rows[0][1]
                found_headers = matched_rows[0][2]
                # Figure out which ID is missing
                missing_ids = []
                for rid in requested_ids:
                    rid_lower = rid.strip().lower()
                    is_found = False
                    for val in found_row.values():
                        if val is not None and rid_lower == str(val).strip().lower():
                            is_found = True
                            break
                    if not is_found:
                        missing_ids.append(rid)
                missing = missing_ids[0] if missing_ids else requested_ids[-1]
                # Build partial result with found student's data
                parts = []
                for h in found_headers:
                    cell = found_row.get(h)
                    if cell is not None and str(cell).strip():
                        parts.append(f"  - {h}: {cell}")
                return (
                    f"PARTIAL VERIFICATION: Cannot fully compare because '{missing}' "
                    f"does not exist in the uploaded data.\n\n"
                    f"Found data for {found_name}:\n" + "\n".join(parts) + "\n\n"
                    f"The student/ID '{missing}' was not found in any of the uploaded documents. "
                    f"This comparison is only partially verified."
                )
            return None

        # Use first two matched students
        (name1, row1, headers1, tkey1) = matched_rows[0]
        (name2, row2, headers2, tkey2) = matched_rows[1]

        # Check if a specific column is asked for comparison
        asked_cols = self._find_asked_columns(query, headers1, tkey1)

        if asked_cols:
            # Compare specific columns
            lines = [f"Comparison between {name1} and {name2}:\n"]
            for col in asked_cols:
                val1 = row1.get(col)
                val2 = row2.get(col)
                v1_f = self._to_float(val1)
                v2_f = self._to_float(val2)
                lines.append(f"  {col}:")
                lines.append(f"    {name1}: {val1}")
                lines.append(f"    {name2}: {val2}")
                if v1_f is not None and v2_f is not None:
                    diff = v1_f - v2_f
                    if diff > 0:
                        lines.append(f"    → {name1} is higher by {abs(diff):.2f}")
                    elif diff < 0:
                        lines.append(f"    → {name2} is higher by {abs(diff):.2f}")
                    else:
                        lines.append(f"    → Both are equal")
            return "\n".join(lines)
        else:
            # Compare all numeric columns
            lines = [f"Comparison between {name1} and {name2}:\n"]
            wins1, wins2 = 0, 0
            for col in headers1:
                v1 = self._to_float(row1.get(col))
                v2 = self._to_float(row2.get(col))
                if v1 is None or v2 is None:
                    continue
                diff = v1 - v2
                marker = ""
                if diff > 0:
                    marker = f" ✓ (+{diff:.1f})"
                    wins1 += 1
                elif diff < 0:
                    marker = f" ✗ ({diff:.1f})"
                    wins2 += 1
                lines.append(f"  {col}: {v1} vs {v2}{marker}")

            lines.append(f"\nSummary: {name1} leads in {wins1} subjects, {name2} leads in {wins2} subjects.")
            total1 = self._to_float(row1.get('TOTAL'))
            total2 = self._to_float(row2.get('TOTAL'))
            pct1 = self._to_float(row1.get('%'))
            pct2 = self._to_float(row2.get('%'))
            if pct1 is not None and pct2 is not None:
                if pct1 > pct2:
                    lines.append(f"Overall: {name1} has higher attendance ({pct1}% vs {pct2}%).")
                elif pct2 > pct1:
                    lines.append(f"Overall: {name2} has higher attendance ({pct2}% vs {pct1}%).")
                else:
                    lines.append(f"Overall: Both have the same attendance percentage ({pct1}%).")
            return "\n".join(lines)

    # Words to strip when extracting a potential name from a query
    _STOP_WORDS = {
        'what', 'is', 'the', 'of', 'tell', 'me', 'about', 'who', 'how',
        'many', 'much', 'give', 'show', 'get', 'find', 'details', 'detail',
        'info', 'information', 'attendance', 'marks', 'score', 'total',
        'percentage', 'classes', 'attended', 'for', 'by', 'a', 'an', 'and',
        'in', 'to', 'does', 'did', 'has', 'have', 'had', 'can', 'do',
        'please', 'sir', 'student', 'roll', 'number', 'name',
    }

    def _check_hallucinated_id(self, query: str) -> Optional[str]:
        """If the query mentions an ID / roll number / name that doesn't exist
        in any table, return a hallucination warning."""

        # Collect all known IDs and names from every table
        known_values = set()
        known_names = []  # list of (lowercase_name, original_name)
        for tkey, rows in self.tables.items():
            for row in rows:
                for val in row.values():
                    if val is not None:
                        val_str = str(val).strip()
                        known_values.add(val_str.lower())
                        # Collect all text values as potential names
                        if isinstance(val, str) and len(val_str) > 2 and self._to_float(val) is None:
                            known_names.append((val_str.lower(), val_str))

        # Look for ID-like patterns in the query
        candidates = []
        for pattern in (self._ID_PATTERN, self._GENERIC_ID):
            candidates.extend(pattern.findall(query))

        # Also check for quoted or capitalized multi-word names
        name_matches = re.findall(r'\b([A-Z][A-Z ]{4,})\b', query)
        candidates.extend(name_matches)

        for candidate in candidates:
            c_lower = candidate.strip().lower()
            if c_lower and c_lower not in known_values:
                return (
                    f"HALLUCINATION DETECTED: '{candidate}' does not exist in the uploaded data. "
                    f"This identifier was not found in any of the loaded documents. "
                    f"The information about '{candidate}' cannot be verified and is likely fabricated."
                )

        # ── Extract a potential name from the query (even lowercase) ────────
        # Strip stop words and see if what remains looks like a person's name
        q_words = re.findall(r'[a-zA-Z]+', query)
        name_words = [w for w in q_words if w.lower() not in self._STOP_WORDS and len(w) > 1]
        extracted_name = " ".join(name_words).strip()

        if len(name_words) >= 1 and extracted_name:
            extracted_lower = extracted_name.lower()
            # Only match if the extracted name is an EXACT full match of a known name
            for known_lower, known_original in known_names:
                if extracted_lower == known_lower:
                    return None  # Exact full name match, not hallucinated

            # Name was extracted but no exact match found
            return (
                f"HALLUCINATION DETECTED: '{extracted_name}' does not exist in the uploaded data. "
                f"No matching student or record was found in the uploaded documents. "
                f"Please use the full name exactly as it appears in the data."
            )

        return None

    def _find_asked_columns(self, query: str, headers: List[str], table_key: str) -> List[str]:
        """Detect which columns the user is asking about in a lookup query.

        Returns a list of matching column names, or empty list if the query
        is generic (e.g. 'tell me about X').
        """
        q_lower = query.lower()
        q_words = set(re.findall(r'\w+', q_lower))
        q_stems = {self._stem(w) for w in q_words if len(w) > 2}

        # If the query is generic ("tell me about X", "details of X"), return empty
        generic_patterns = [r'\btell\b.*\babout\b', r'\bdetails?\b.*\bof\b',
                            r'\binfo\b.*\babout\b', r'\ball\b.*\bdetails?\b',
                            r'\bshow\b.*\bdata\b', r'\bfull\b.*\bdata\b']
        if any(re.search(p, q_lower) for p in generic_patterns):
            return []

        # Skip these generic words that don't refer to columns
        skip_words = {'what', 'how', 'many', 'the', 'who', 'which', 'tell',
                      'about', 'give', 'show', 'get', 'find', 'is', 'are',
                      'was', 'were', 'has', 'have', 'had', 'does', 'did',
                      'classes', 'attended', 'scored', 'marks', 'score',
                      'value', 'number', 'much', 'detail', 'info',
                      'student', 'name', 'roll', 'sir', 'please', 'of', 'by'}

        # First: check if a full column name appears verbatim in the query
        # e.g. "fml lab" in "How many FML LAB classes attended by X?"
        # Sort by length descending so "FML LAB" matches before "FML"
        exact_matches = []
        for col in sorted(headers, key=lambda c: len(c), reverse=True):
            col_lower = col.lower().strip()
            # Check aliases first (even for single-char columns like '%')
            aliases = set()
            for alias_key, alias_set in self.COLUMN_ALIASES.items():
                if col_lower == alias_key or col_lower in alias_set:
                    aliases = alias_set
                    break
            if aliases and (q_words & aliases):
                exact_matches.append(col)
                continue
            if len(col_lower) < 2:
                continue
            # For short column names (<=3 chars like "SE", "OS"), use word boundary
            # to avoid matching inside other words ("se" in "classes")
            if len(col_lower) <= 3:
                if re.search(r'\b' + re.escape(col_lower) + r'\b', q_lower):
                    exact_matches.append(col)
            else:
                # Longer names: verbatim substring is fine
                if col_lower in q_lower:
                    exact_matches.append(col)

        if exact_matches:
            # Filter out columns whose names are substrings of already-matched longer names
            # e.g. if "FML LAB" matched, don't also return "FML"
            filtered = []
            for col in exact_matches:
                cl = col.lower().strip()
                is_substring = any(
                    cl != other.lower().strip() and cl in other.lower().strip()
                    for other in exact_matches
                )
                if not is_substring:
                    filtered.append(col)
            return filtered

        # Fallback: stem/substring matching for partial names
        matched = []
        for col in headers:
            col_lower = col.lower().strip()
            col_words = set(re.findall(r'\w+', col_lower))
            col_stems = {self._stem(w) for w in col_words}

            if not col_stems:
                continue

            stem_hits = len(q_stems & col_stems)
            sub_hits = sum(
                1 for qw in q_words - skip_words
                if len(qw) > 1 and any(
                    (qw == cw or (len(qw) > 2 and len(cw) > 2 and (qw in cw or cw in qw)))
                    for cw in col_words
                )
            )

            if stem_hits > 0 or sub_hits > 0:
                matched.append(col)

        return matched

    # ── Filter / Conditional Queries ─────────────────────────────────────────
    def _try_filter_query(self, query: str) -> Optional[str]:
        """Answer queries like 'how many students have attendance > 80%' or
        'list students with percentage above 90'."""
        q_lower = query.lower()

        # Detect a filter condition
        filter_op = None
        threshold = None
        threshold2 = None  # for 'between'

        for pattern, op in FILTER_PATTERNS:
            m = re.search(pattern, q_lower)
            if m:
                filter_op = op
                if op == "between":
                    threshold = float(m.group(1))
                    threshold2 = float(m.group(2))
                else:
                    threshold = float(m.group(2))
                break

        if filter_op is None:
            return None

        # Find the column to filter on
        table_key, column = self._match_column(query, "max")
        if table_key is None or column is None:
            # Try first table, % column
            table_key = next(iter(self.tables), None)
            if table_key is None:
                return None
            # Look for a % or percentage column
            for h in self.headers[table_key]:
                if h.strip() in ('%', 'Percentage', 'percentage', 'Attendance'):
                    column = h
                    break
            if column is None:
                return None

        rows = self.tables[table_key]
        label_col = self._find_label_column(rows)

        # Apply the filter
        matching = []
        for r in rows:
            val = self._to_float(r.get(column))
            if val is None:
                continue
            label = str(r.get(label_col, "?")).strip() if label_col else "?"

            if filter_op == "gte" and val >= threshold:
                matching.append((label, val))
            elif filter_op == "lte" and val <= threshold:
                matching.append((label, val))
            elif filter_op == "eq" and abs(val - threshold) < 0.01:
                matching.append((label, val))
            elif filter_op == "between" and threshold <= val <= threshold2:
                matching.append((label, val))

        matching.sort(key=lambda x: x[1], reverse=True)
        col_clean = column.strip()

        # Detect if query asks "how many" (count) or "list/who" (list names)
        wants_count = bool(re.search(r"(how many|count|number of)", q_lower))
        op_label = {
            "gte": f"greater than or equal to {threshold}",
            "lte": f"less than or equal to {threshold}",
            "eq": f"equal to {threshold}",
            "between": f"between {threshold} and {threshold2}",
        }[filter_op]

        if wants_count:
            answer = f"{len(matching)} students have {col_clean} {op_label}."
            if matching and len(matching) <= 20:
                names = ", ".join(f"{lbl} ({v})" for lbl, v in matching[:10])
                answer += f"\n\nThey are: {names}"
                if len(matching) > 10:
                    answer += f" ... and {len(matching) - 10} more."
            return answer
        else:
            # List them
            if not matching:
                return f"No students found with {col_clean} {op_label}."
            lines = [f"Students with {col_clean} {op_label} ({len(matching)} found):"]
            for i, (lbl, v) in enumerate(matching[:20], 1):
                lines.append(f"  {i}. {lbl} — {v}")
            if len(matching) > 20:
                lines.append(f"  ... and {len(matching) - 20} more.")
            return "\n".join(lines)

    # ── Internal helpers ─────────────────────────────────────────────────────
    @staticmethod
    def _extract_claimed_value(query: str) -> Optional[float]:
        """Extract a numeric value the user is claiming/asserting in the query.
        e.g. 'is 22PA1A0501 has attendance percentage of 90%' → 90.0
             'does X have 85 marks' → 85.0
        Only triggers for verification-style queries (is/does/has/did/correct/true).
        """
        q_lower = query.lower()
        # Only look for claimed values in verification-style queries
        verification_words = ('is ', 'does ', 'has ', 'did ', 'had ', 'correct', 'true', 'right')
        if not any(q_lower.startswith(w) or w in q_lower for w in verification_words):
            return None
        # Extract numbers from the query (skip roll-number-like patterns)
        numbers = re.findall(r'(?<!\w)(\d+\.?\d*)%?(?!\w*[A-Za-z])', query)
        # Filter out roll-number-like values (long alphanumeric codes)
        roll_pattern = re.compile(r'\d{2}[A-Za-z]{2}\d[A-Za-z]\d{4}')
        roll_numbers = roll_pattern.findall(query)
        roll_digits = set()
        for rn in roll_numbers:
            roll_digits.update(re.findall(r'\d+', rn))
        # Return the last number that isn't part of a roll number
        for num_str in reversed(numbers):
            if num_str not in roll_digits:
                try:
                    return float(num_str)
                except ValueError:
                    continue
        return None

    @staticmethod
    def _stem(word: str) -> str:
        """Cheap suffix stripping so 'students' matches 'student' etc."""
        w = word.lower()
        for suffix in ("ing", "tion", "ness", "ment", "ies", "es", "ed", "ly", "s"):
            if len(w) > len(suffix) + 2 and w.endswith(suffix):
                return w[: -len(suffix)]
        return w

    def _detect_operation(self, query: str) -> Optional[str]:
        q = query.lower()
        for pattern, op in AGGREGATE_PATTERNS:
            if re.search(pattern, q):
                return op
        return None

    # Map short / symbolic column names to query-friendly aliases
    COLUMN_ALIASES = {
        '%': {'percentage', 'percent', 'attendance', 'rate'},
        'total': {'total', 'overall', 'sum', 'aggregate'},
        'p&s': {'p&s', 'ps', 'p and s', 'probability', 'p s'},
    }

    def _match_column(self, query: str, op: str = None):
        """Find which table + column the query is about.

        Uses stemming, substring matching, and alias expansion so that
        e.g. 'students' matches 'Student Name', 'attendance percentage'
        matches the '%' column, etc.
        """
        q_lower = query.lower()
        q_stems = {self._stem(w) for w in re.findall(r'\w+', q_lower) if len(w) > 2}
        q_words = set(re.findall(r'\w+', q_lower))

        best_score = 0.0
        best_table = None
        best_col = None

        for tkey, headers in self.headers.items():
            for col in headers:
                col_lower = col.lower().strip()
                col_words = set(re.findall(r'\w+', col_lower))
                col_stems = {self._stem(w) for w in col_words}

                # --- Check aliases for short/symbolic column names ---
                aliases = set()
                for alias_key, alias_set in self.COLUMN_ALIASES.items():
                    if col_lower == alias_key or col_lower in alias_set:
                        aliases = alias_set
                        break

                alias_hits = len(q_words & aliases) if aliases else 0

                if alias_hits > 0:
                    # Strong match via alias
                    score = 0.9 + alias_hits * 0.05
                elif not col_stems:
                    continue
                else:
                    # Method 1: stem-based overlap
                    stem_overlap = len(q_stems & col_stems)
                    score1 = stem_overlap / len(col_stems) if col_stems else 0

                    # Method 2: substring match (skip 1-char stems to avoid false positives)
                    sub_hits = 0
                    for qw in q_stems:
                        if any(
                            (qw in cw or cw in qw) and len(cw) > 1 and len(qw) > 1
                            for cw in col_stems
                        ):
                            sub_hits += 1
                    score2 = sub_hits / len(col_stems) if col_stems else 0

                    score = max(score1, score2)

                # For numeric aggregations, prefer numeric columns
                if op in ("max", "min", "avg", "sum", "sort") and score > 0:
                    rows = self.tables[tkey]
                    sample_val = rows[0].get(col) if rows else None
                    if self._to_float(sample_val) is not None:
                        score += 0.1  # small boost for numeric cols

                if score > best_score:
                    best_score = score
                    best_table = tkey
                    best_col = col

        if best_score < 0.25:
            return None, None
        return best_table, best_col

    def _to_float(self, val) -> Optional[float]:
        """Try to parse a cell value as float."""
        if val is None:
            return None
        s = str(val).strip().replace("%", "").replace(",", "").replace("$", "")
        try:
            return float(s)
        except (ValueError, TypeError):
            return None

    def _find_label_column(self, rows: List[Dict]) -> Optional[str]:
        """Find the column that likely contains names/labels."""
        if not rows:
            return None
        # Prefer columns with 'name' in the header
        for col in rows[0]:
            if 'name' in col.lower():
                return col
        # Fallback: first column whose values are mostly non-numeric strings
        for col in rows[0]:
            non_num = sum(1 for r in rows[:10] if r.get(col) and self._to_float(r[col]) is None)
            if non_num > len(rows[:10]) * 0.5:
                return col
        return list(rows[0].keys())[0]

    def _compute(self, rows: List[Dict], column: Optional[str], op: str, query: str) -> Optional[str]:
        """Run the aggregate and build a natural-language answer."""
        label_col = self._find_label_column(rows)

        # For count, we can work without a numeric column
        if op == "count":
            total = len(rows)
            if column and column != label_col:
                # Count non-empty values in that column
                filled = sum(1 for r in rows if r.get(column) is not None and str(r.get(column)).strip())
                return f"There are {filled} entries with {column} values (out of {total} total rows)."
            return f"There are {total} entries/rows in the data."

        if column is None:
            return None

        # Extract numeric values paired with their labels
        pairs = []
        for r in rows:
            val = self._to_float(r.get(column))
            label = str(r.get(label_col, "?")).strip() if label_col else "?"
            if val is not None:
                pairs.append((label, val))

        if not pairs:
            return None

        col_clean = column.strip()

        if op == "max":
            pairs.sort(key=lambda x: x[1], reverse=True)
            winner = pairs[0]
            answer = f"{winner[0]} has the highest {col_clean} with a value of {winner[1]}."
            if len(pairs) > 1:
                answer += f" Followed by {pairs[1][0]} ({pairs[1][1]})"
                if len(pairs) > 2:
                    answer += f" and {pairs[2][0]} ({pairs[2][1]})"
                answer += "."
            return answer

        if op == "min":
            pairs.sort(key=lambda x: x[1])
            winner = pairs[0]
            answer = f"{winner[0]} has the lowest {col_clean} with a value of {winner[1]}."
            if len(pairs) > 1:
                answer += f" Followed by {pairs[1][0]} ({pairs[1][1]})"
                if len(pairs) > 2:
                    answer += f" and {pairs[2][0]} ({pairs[2][1]})"
                answer += "."
            return answer

        if op == "avg":
            vals = [v for _, v in pairs]
            avg = sum(vals) / len(vals)
            return f"The average {col_clean} is {avg:.2f} (across {len(vals)} entries)."

        if op == "sum":
            total = sum(v for _, v in pairs)
            return f"The total {col_clean} is {total:.2f} (across {len(pairs)} entries)."

        if op == "count":
            return f"There are {len(pairs)} entries with numeric {col_clean} values."

        if op == "sort":
            pairs.sort(key=lambda x: x[1], reverse=True)
            lines = [f"Ranking by {col_clean} (highest to lowest):"]
            for i, (lbl, val) in enumerate(pairs[:15], 1):
                lines.append(f"  {i}. {lbl} — {val}")
            if len(pairs) > 15:
                lines.append(f"  ... and {len(pairs) - 15} more.")
            return "\n".join(lines)

        return None

    # ── LLM-powered Data Analysis ────────────────────────────────────────────
    def _try_llm_analysis(self, query: str) -> Optional[str]:
        """Use Groq LLM to analyze structured data for complex questions
        that the pattern-based methods can't handle."""
        if Groq is None:
            return None

        from config.settings import GROQ_API_KEY, LLM_MODEL
        if not GROQ_API_KEY:
            return None

        # Build a compact data summary for the LLM
        data_context = self._build_data_context()
        if not data_context:
            return None

        prompt = f"""You are a data analyst. Answer the following question using ONLY the data provided below.
Be precise and use actual numbers from the data. If the answer cannot be determined from the data, say so.
Do not include file paths, source references, or [Source: ...] tags.
Give a clear, natural response.

DATA:
{data_context}

QUESTION: {query}

ANSWER:"""

        try:
            client = Groq(api_key=GROQ_API_KEY)
            response = client.chat.completions.create(
                model=LLM_MODEL,
                messages=[
                    {"role": "system", "content": "You are a precise data analyst. Answer only from the given data. Be concise and accurate."},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=1000,
                temperature=0.1
            )
            answer = response.choices[0].message.content.strip()
            if answer:
                return answer
        except Exception as e:
            print(f"LLM analysis error: {e}")

        return None

    def _build_data_context(self, max_rows: int = 80) -> str:
        """Convert stored tables into a compact text format for LLM context."""
        parts = []
        for tkey, rows in self.tables.items():
            headers = self.headers.get(tkey, [])
            if not rows:
                continue

            parts.append(f"Table: {tkey}")
            parts.append(f"Columns: {', '.join(headers)}")
            parts.append(f"Total rows: {len(rows)}")

            # Include data as CSV-like format (compact)
            parts.append("Data:")
            parts.append(" | ".join(headers))
            for r in rows[:max_rows]:
                vals = [str(r.get(h, "")) for h in headers]
                parts.append(" | ".join(vals))

            if len(rows) > max_rows:
                parts.append(f"... ({len(rows) - max_rows} more rows)")
            parts.append("")

        return "\n".join(parts)
